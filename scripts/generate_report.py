from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


# ---------------------------------------------------------
# Paths
# ---------------------------------------------------------

PROJECT_ROOT = Path(__file__).resolve().parent.parent

REPORTS_DIR = PROJECT_ROOT / "reports"
REPORTS_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = REPORTS_DIR / "Marriott_Financial_Audit_Analytical_Review.docx"


# ---------------------------------------------------------
# Document setup
# ---------------------------------------------------------

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10.5)


# ---------------------------------------------------------
# Helper functions
# ---------------------------------------------------------

def add_heading(text, level=1):
    doc.add_heading(text, level=level)


def add_bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(headers, rows):
    table = doc.add_table(
        rows=1,
        cols=len(headers)
    )

    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header
    for cell, header in zip(table.rows[0].cells, headers):
        cell.text = str(header)

        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True

    # Rows
    for row in rows:
        cells = table.add_row().cells

        for cell, value in zip(cells, row):
            cell.text = str(value)

    return table


# ---------------------------------------------------------
# Title
# ---------------------------------------------------------

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = title.add_run(
    "Marriott International, Inc.\n"
    "Financial & Audit Analytical Review"
)

run.bold = True
run.font.size = Pt(20)


subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = subtitle.add_run(
    "Initial analytical report — FY2015–FY2025"
)

run.italic = True
run.font.size = Pt(11)


# ---------------------------------------------------------
# 1. Objective and scope
# ---------------------------------------------------------

add_heading("1. Objective and scope")

doc.add_paragraph(
    "The objective of this project is to develop a reproducible financial "
    "and audit-oriented analytical review of Marriott International, Inc. "
    "using financial information obtained from SEC filings and XBRL data."
)

doc.add_paragraph(
    "The analysis combines Python-based data preparation, financial-ratio "
    "analysis and audit-oriented analytical procedures."
)

doc.add_paragraph(
    "The current analysis covers fiscal years 2015–2025 and focuses on "
    "revenue, receivables, operating cash flow, payables, working-capital "
    "effects, profitability and cash conversion."
)


# ---------------------------------------------------------
# 2. Data source and methodology
# ---------------------------------------------------------

add_heading("2. Data source and methodology")

doc.add_paragraph(
    "The principal source is the U.S. Securities and Exchange Commission "
    "(SEC) EDGAR/XBRL financial data for Marriott International, Inc. "
    "(CIK 0001048286)."
)

doc.add_paragraph(
    "Annual facts were selected and cleaned to avoid duplicate annual "
    "observations appearing in successive SEC filings."
)

doc.add_paragraph(
    "Python and pandas are used as the analytical engine. Calculations "
    "are kept reproducible in the project dataframe."
)

add_bullets([
    "Growth rates are calculated from consecutive annual observations.",
    "Operating cash flow is used to evaluate cash conversion.",
    "EBIT is calculated as profit before tax plus interest expense.",
    "EBITA is calculated as EBIT plus amortization where reliable data is available.",
    "EBITDA is calculated as EBIT plus depreciation plus amortization where both components are available.",
    "Missing amortization data is retained rather than estimated."
])


# ---------------------------------------------------------
# 3. Working-capital analysis
# ---------------------------------------------------------

add_heading("3. Working-capital cash effect")

doc.add_paragraph(
    "A simplified working-capital cash effect was calculated from changes "
    "in receivables and payables:"
)

formula = doc.add_paragraph()
formula.alignment = WD_ALIGN_PARAGRAPH.CENTER

run = formula.add_run(
    "Working-capital cash effect = change in payables − change in receivables"
)

run.bold = True

doc.add_paragraph(
    "A negative amount represents a simplified use of cash, while a positive "
    "amount represents a simplified source of cash. This measure includes "
    "only receivables and payables and therefore should not be interpreted "
    "as the complete operating-working-capital adjustment in the statement "
    "of cash flows."
)

working_capital_rows = [
    [2016, "$(498)M", "-31.48%", "-2.92%"],
    [2017, "$(200)M", "-8.21%", "-0.87%"],
    [2018, "$(158)M", "-6.70%", "-0.76%"],
    [2019, "$(309)M", "-18.34%", "-1.47%"],
    [2020, "$434M", "26.48%", "4.11%"],
    [2021, "$(15)M", "-1.27%", "-0.11%"],
    [2022, "$(569)M", "-24.08%", "-2.74%"],
    [2023, "$(149)M", "-4.70%", "-0.63%"],
    [2024, "$(58)M", "-2.11%", "-0.23%"],
    [2025, "$(63)M", "-1.96%", "-0.24%"],
]

add_table(
    ["Year", "WC cash effect", "% of CFO", "% of revenue"],
    working_capital_rows
)


# ---------------------------------------------------------
# 4. Audit-oriented screening
# ---------------------------------------------------------

add_heading("4. Audit-oriented analytical screening")

doc.add_paragraph(
    "A 15% threshold was used as an analytical screening device in selected "
    "procedures. This threshold is not an audit materiality determination. "
    "Actual audit materiality depends on the engagement, benchmark, "
    "qualitative factors and professional judgment."
)

screening_rows = [
    [
        2016,
        "Receivables grew materially faster than revenue",
        "Large relative to CFO",
        "Large relative to PBT",
        "Starwood acquisition year",
    ],
    [
        2019,
        "Receivables grew materially faster than revenue",
        "Large relative to CFO",
        "Large relative to PBT",
        "No major exceptional context identified",
    ],
    [
        2020,
        "Receivables grew materially faster than revenue",
        "Large relative to CFO",
        "Not applicable",
        "COVID-19 disruption year",
    ],
    [
        2022,
        "Revenue grew materially faster than receivables",
        "Large relative to CFO",
        "Large relative to PBT",
        "No major exceptional context identified",
    ],
]

add_table(
    [
        "Year",
        "Growth signal",
        "Cash-flow signal",
        "Profit signal",
        "Context",
    ],
    screening_rows
)

doc.add_paragraph(
    "These flags identify periods that warrant further investigation. "
    "They do not establish that a material misstatement exists. Further "
    "audit evidence would be required to reach such a conclusion."
)


# ---------------------------------------------------------
# 5. Profitability
# ---------------------------------------------------------

add_heading("5. Profitability: EBIT and EBITDA")

doc.add_paragraph(
    "EBIT and EBITDA were constructed from available SEC/XBRL information. "
    "Because reliable annual amortization data is not available for every "
    "year, EBITDA is deliberately left unavailable for unsupported periods."
)

profitability_rows = [
    [2015, "$1,422M", "$1,619M", "9.82%", "11.18%"],
    [2016, "$1,418M", "$1,662M", "8.31%", "9.74%"],
    [2017, "$3,124M", "$3,540M", "13.65%", "15.46%"],
    [2018, "$2,685M", "—", "12.93%", "—"],
    [2019, "$1,993M", "—", "9.50%", "—"],
    [2020, "$(21)M", "—", "-0.20%", "—"],
    [2021, "$1,600M", "—", "11.55%", "—"],
    [2022, "$3,517M", "$3,828M", "16.93%", "18.43%"],
    [2023, "$3,943M", "$4,291M", "16.63%", "18.10%"],
    [2024, "$3,846M", "$4,229M", "15.32%", "16.85%"],
    [2025, "$4,203M", "$4,661M", "16.05%", "17.80%"],
]

add_table(
    ["Year", "EBIT", "EBITDA", "EBIT margin", "EBITDA margin"],
    profitability_rows
)

doc.add_paragraph(
    "Marriott's disclosed non-GAAP Adjusted EBITDA reconciliations were "
    "reviewed as a cross-check. Those reconciliations contain broader "
    "management-defined D&A-related components and are therefore not "
    "substituted for the standardized EBITDA calculation used in this project."
)


# ---------------------------------------------------------
# 6. Cash conversion
# ---------------------------------------------------------

add_heading("6. Profitability and cash conversion")

doc.add_paragraph(
    "EBIT margin measures operating profitability relative to revenue. "
    "CFO/EBIT measures operating cash flow relative to EBIT. The latter "
    "is used as an analytical cash-conversion indicator."
)

cash_conversion_rows = [
    [2015, "9.82%", "11.18%", "100.56%"],
    [2016, "8.31%", "9.74%", "111.57%"],
    [2017, "13.65%", "15.46%", "77.98%"],
    [2018, "12.93%", "—", "87.78%"],
    [2019, "9.50%", "—", "84.55%"],
    [2020, "-0.20%", "—", "—"],
    [2021, "11.55%", "—", "73.56%"],
    [2022, "16.93%", "18.43%", "67.19%"],
    [2023, "16.63%", "18.10%", "80.40%"],
    [2024, "15.32%", "16.85%", "71.48%"],
    [2025, "16.05%", "17.80%", "76.42%"],
]

add_table(
    ["Year", "EBIT margin", "EBITDA margin", "CFO / EBIT"],
    cash_conversion_rows
)

# ---------------------------------------------------------
# 6.1 Net income to operating cash flow
# ---------------------------------------------------------

add_heading("6.1 Net income to operating cash flow", level=2)

doc.add_paragraph(
    "Because CFO begins with net income under the indirect method, "
    "CFO/net income provides a complementary cash-conversion measure "
    "to CFO/EBIT. This distinction is important because EBIT is measured "
    "before interest and income taxes, whereas CFO starts from net income."
)

net_income_cfo_rows = [
    [2022, "$2,358M", "$2,363M", "$5M", "100.21%"],
    [2024, "$2,375M", "$2,749M", "$374M", "115.75%"],
    [2025, "$2,601M", "$3,212M", "$611M", "123.49%"],
]

add_table(
    [
        "Year",
        "Net income",
        "CFO",
        "CFO − net income",
        "CFO / net income",
    ],
    net_income_cfo_rows
)

doc.add_paragraph(
    "The results materially change the interpretation of the post-COVID "
    "cash-conversion picture. In 2024 and 2025, operating cash flow was "
    "substantially higher than net income, indicating strong cash conversion "
    "relative to reported earnings. The lower CFO/EBIT ratios in those years "
    "should therefore not be interpreted in isolation as evidence of weak "
    "cash conversion."
)

add_heading("6.2 Known non-cash adjustments", level=2)

doc.add_paragraph(
    "To investigate the difference between net income and operating cash "
    "flow, selected non-cash adjustments were examined. These include "
    "depreciation, amortization where reliably available, and share-based "
    "compensation."
)

non_cash_rows = [
    [2022, "$5M", "$503M", "$(498)M"],
    [2024, "$374M", "$620M", "$(246)M"],
    [2025, "$611M", "$694M", "$(83)M"],
]

add_table(
    [
        "Year",
        "CFO − net income",
        "Known non-cash adjustments",
        "Residual after selected adjustments",
    ],
    non_cash_rows
)

doc.add_paragraph(
    "The residual represents the difference remaining after subtracting "
    "the selected non-cash adjustments from the CFO increase over net income. "
    "It should not be interpreted as a complete working-capital adjustment, "
    "because other operating and reconciliation items are also included."
)

doc.add_paragraph(
    "In 2022, $503 million of identified non-cash adjustments were almost "
    "entirely offset by other operating effects, leaving only a $5 million "
    "net difference between CFO and net income. In 2024 and 2025, larger "
    "portions of the identified non-cash adjustments remained reflected in "
    "operating cash flow."
)

add_heading("6.3 Audit interpretation of cash conversion", level=2)

doc.add_paragraph(
    "The analysis indicates that the CFO/EBIT ratio should be treated as "
    "an analytical screening indicator rather than as a standalone measure "
    "of cash conversion quality."
)

doc.add_paragraph(
    "The 2022 result remains particularly relevant because strong "
    "profitability coincided with CFO only slightly above net income and "
    "with a substantial offset between identified non-cash additions and "
    "other operating effects."
)

doc.add_paragraph(
    "The simplified receivables/payables analysis identified a $569 million "
    "use of cash in 2022. This is broadly consistent in magnitude with the "
    "approximately $498 million residual identified after selected non-cash "
    "adjustments, although the two measures are not identical and should "
    "not be treated as a formal reconciliation."
)

doc.add_paragraph(
    "For 2024 and 2025, CFO exceeded net income by $374 million and "
    "$611 million respectively. The simplified receivables/payables effect "
    "was comparatively small in those years, suggesting that other operating "
    "items explain a larger proportion of the difference."
)



# ---------------------------------------------------------
# 7. Key analytical observation
# ---------------------------------------------------------

add_heading("7. Key analytical observation")

doc.add_paragraph(
    "The most notable period is 2022. Marriott recorded an EBIT margin "
    "of 16.93% and a supported EBITDA margin of 18.43%, while CFO/EBIT "
    "was only 67.19%."
)

doc.add_paragraph(
    "Further analysis shows that CFO was nevertheless broadly in line with "
    "net income in 2022: CFO was $2.363 billion compared with net income "
    "of $2.358 billion. The simplified receivables/payables analysis "
    "identified a $569 million use of cash, while selected known non-cash "
    "adjustments amounted to $503 million."
)

doc.add_paragraph(
    "This indicates that the 2022 cash-flow profile deserves further "
    "investigation, particularly because substantial non-cash additions "
    "were largely offset by other operating effects."
)

doc.add_paragraph(
    "In contrast, 2024 and 2025 show stronger cash conversion relative to "
    "net income. CFO exceeded net income by $374 million in 2024 and "
    "$611 million in 2025. Therefore, the lower CFO/EBIT ratios in those "
    "years should not, by themselves, be interpreted as evidence of weak "
    "operating cash generation."
)

doc.add_paragraph(
    "The overall analytical conclusion is therefore more nuanced: 2022 "
    "represents the strongest period for further investigation, whereas "
    "2024–2025 demonstrate strong conversion of reported net income into "
    "operating cash flow."
)


# ---------------------------------------------------------
# 8. Limitations
# ---------------------------------------------------------

add_heading("8. Audit interpretation and limitations")

add_bullets([
    "Analytical flags identify relationships that warrant investigation; they do not prove a misstatement.",
    "2020 represents an exceptional COVID-19 disruption period.",
    "2016 should be interpreted in the context of Marriott's Starwood acquisition.",
    "The simplified working-capital measure includes only receivables and payables.",
    "CFO/EBIT is excluded when EBIT is negative.",
    "CFO/EBIT is an analytical screening ratio rather than a direct reconciliation between operating profit and operating cash flow.",
    "CFO/net income provides a complementary measure because the indirect cash-flow statement begins with net income.",
    "The selected non-cash adjustments do not constitute a complete CFO reconciliation.",
    "EBITDA is unavailable for 2018–2021 because reliable annual amortization data was not available.",
    "Marriott's Adjusted EBITDA is a management-defined non-GAAP measure.",
    "Actual audit materiality and performance materiality require engagement-specific professional judgment.",
])




# ---------------------------------------------------------
# 9. Current conclusion
# ---------------------------------------------------------

add_heading("9. Current conclusion")

doc.add_paragraph(
    "The analysis to date shows a significant improvement in Marriott's "
    "operating profitability after the COVID-19 disruption, with EBIT "
    "margins reaching approximately 15–17% during 2022–2025."
)

doc.add_paragraph(
    "The CFO/EBIT ratio provides a useful analytical screening indicator, "
    "but it should not be interpreted independently because EBIT is measured "
    "before interest and income taxes, whereas CFO begins with net income."
)

doc.add_paragraph(
    "The complementary CFO/net income analysis shows that operating cash "
    "flow exceeded net income substantially in both 2024 and 2025. These "
    "years therefore do not indicate weak cash conversion relative to "
    "reported earnings."
)

doc.add_paragraph(
    "2022 remains the strongest candidate for further investigation. "
    "The year combines strong profitability with CFO only marginally above "
    "net income, a $569 million simplified receivables/payables use of cash, "
    "and $503 million of identified non-cash adjustments that were largely "
    "offset by other operating effects."
)

doc.add_paragraph(
    "From an audit-analytical perspective, these relationships represent "
    "risk indicators and areas for further investigation rather than "
    "evidence of a material misstatement."
)


# ---------------------------------------------------------
# Save
# ---------------------------------------------------------

doc.save(OUTPUT_FILE)

print(f"Report created successfully:")
print(OUTPUT_FILE)